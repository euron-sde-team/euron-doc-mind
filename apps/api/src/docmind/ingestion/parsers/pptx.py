"""PowerPoint: one section per slide, with the title, body text, tables and speaker notes.

Speaker notes often hold the substance a slide only hints at, so they are indexed with the
slide they belong to and cited as that slide.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from docmind.core.text import normalize
from docmind.ingestion.archive import check_ooxml
from docmind.ingestion.types import Block, ParsedDocument, ParsedSection, ParseError, ParseLimits
from docmind.models.enums import SectionKind


def _shapes(shapes: Any) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _shape_blocks(shape: Any) -> list[Block]:
    blocks: list[Block] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = normalize("".join(run.text for run in paragraph.runs) or paragraph.text or "")
            if text:
                prefix = (
                    "\u2022 " if paragraph.level > 0 or len(shape.text_frame.paragraphs) > 1 else ""
                )
                blocks.append(Block(prefix + text))
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[normalize(cell.text) for cell in row.cells] for row in shape.table.rows]
        if rows:
            header, *body = rows
            for row in body or [header]:
                pairs = [f"{h}: {v}" for h, v in zip(header, row, strict=False) if v and h]
                blocks.append(Block("; ".join(pairs) if pairs and body else " | ".join(row)))
    return blocks


def parse_pptx(data: bytes, limits: ParseLimits) -> ParsedDocument:
    check_ooxml(
        data,
        required_prefix="ppt/",
        label="PowerPoint deck",
        max_uncompressed=limits.max_uncompressed_bytes,
    )
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("This PowerPoint deck could not be read.") from exc

    sections: list[ParsedSection] = []
    empty = 0
    for number, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        title = (
            normalize(title_shape.text)
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        blocks = [Block(title, heading_level=1)] if title else []
        for shape in _shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            blocks.extend(_shape_blocks(shape))
        if slide.has_notes_slide:
            notes = normalize(slide.notes_slide.notes_text_frame.text or "")
            if notes:
                blocks.append(Block(f"Speaker notes: {notes}"))
        if not any(block.heading_level is None for block in blocks) and not title:
            empty += 1
            continue
        label = f"Slide {number}: {title}" if title else f"Slide {number}"
        sections.append(ParsedSection(SectionKind.SLIDE, label[:300], blocks))

    if not sections:
        raise ParseError("This deck has no text on any slide.")
    deck_title = (deck.core_properties.title or "").strip() or None
    warnings = [f"{empty} slide(s) had no text and were skipped."] if empty else []
    return ParsedDocument(title=deck_title, sections=sections, warnings=warnings)
