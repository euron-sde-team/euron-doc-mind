"""Spreadsheets (XLSX) and CSV: one section per sheet, one block per row.

Rows become "Column: value" records, so a retrieved row reads as a fact ("Department:
Marketing; FY2027 budget: $1,250,000") rather than a bare list of numbers. Excel number
formats are honoured, so percentages and currency survive.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from typing import Any

import openpyxl

from docmind.core.text import normalize
from docmind.ingestion.archive import check_ooxml
from docmind.ingestion.types import Block, ParsedDocument, ParsedSection, ParseError, ParseLimits
from docmind.models.enums import SectionKind


def format_value(value: Any, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float):
        fmt = number_format or "General"
        if "%" in fmt:
            decimals = 1 if "." in fmt else 0
            return f"{value * 100:.{decimals}f}%"
        currency = next(
            (symbol for symbol in ("$", "\u20ac", "\u00a3", "\u20b9") if symbol in fmt), ""
        )
        if currency:
            return f"{currency}{value:,.2f}".replace(".00", "")
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, int):
            return f"{value:,}" if abs(value) >= 10_000 else str(value)
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    return normalize(str(value))


def rows_to_blocks(
    rows: Iterable[Sequence[str]], *, max_rows: int
) -> tuple[list[Block], int, bool]:
    """Header detection plus "Column: value" records. Returns blocks, data rows, truncated."""
    header: list[str] | None = None
    blocks: list[Block] = []
    data_rows = 0
    truncated = False
    for row in rows:
        cells = [cell.strip() for cell in row]
        if not any(cells):
            continue
        if header is None:
            header = [cell or f"Column {i + 1}" for i, cell in enumerate(cells)]
            continue
        if data_rows >= max_rows:
            truncated = True
            break
        pairs = [f"{name}: {value}" for name, value in zip(header, cells, strict=False) if value]
        blocks.append(Block("; ".join(pairs)))
        data_rows += 1
    if header is not None:
        summary = f"Columns: {', '.join(header)}. {data_rows} data row(s)."
        blocks.insert(0, Block(summary))
    return blocks, data_rows, truncated


def parse_xlsx(data: bytes, limits: ParseLimits) -> ParsedDocument:
    check_ooxml(
        data,
        required_prefix="xl/",
        label="Excel workbook",
        max_uncompressed=limits.max_uncompressed_bytes,
    )
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("This Excel workbook could not be read.") from exc

    sections: list[ParsedSection] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = (
                [
                    format_value(cell.value, getattr(cell, "number_format", "General"))
                    for cell in row
                ]
                for row in sheet.iter_rows()
            )
            blocks, _, truncated = rows_to_blocks(rows, max_rows=limits.max_rows_per_sheet)
            if truncated:
                warnings.append(
                    f"Sheet '{sheet.title}' has more than {limits.max_rows_per_sheet:,} rows; "
                    "the rest were not indexed."
                )
            if blocks:
                sections.append(ParsedSection(SectionKind.SHEET, f"Sheet: {sheet.title}", blocks))
        title = (workbook.properties.title or "").strip() or None
    finally:
        workbook.close()

    if not sections:
        raise ParseError("This workbook has no data.")
    return ParsedDocument(title=title, sections=sections, warnings=warnings)


def parse_csv(data: bytes, limits: ParseLimits) -> ParsedDocument:
    text = decode_text(data)
    sample = text[:16_384]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    blocks, _, truncated = rows_to_blocks(reader, max_rows=limits.max_rows_per_sheet)
    if not blocks:
        raise ParseError("This CSV file has no data.")
    warnings = (
        [f"Only the first {limits.max_rows_per_sheet:,} rows were indexed."] if truncated else []
    )
    return ParsedDocument(
        title=None, sections=[ParsedSection(SectionKind.SHEET, "Table", blocks)], warnings=warnings
    )


def decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")
